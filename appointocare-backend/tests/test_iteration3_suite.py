import hmac
import hashlib
import json
import io
import unittest
from datetime import datetime, timedelta
from app import create_app
from app.models import db, Organization, Customer, Appointment, IndustryRecord, FinancialLedgerEntry, AuditLog
from app.utils.jwt_helper import generate_access_token


class Iteration3SuiteTestCase(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.app = create_app()
        cls.client = cls.app.test_client()

    def setUp(self):
        self.ctx = self.app.app_context()
        self.ctx.push()

    def tearDown(self):
        self.ctx.pop()

    def test_customer_export_xlsx_and_formula_injection_defense(self):
        """Validates XLSX export format and sanitization of malicious formula cells."""
        token = generate_access_token(1, "Organization", extra_claims={"organization_id": "1", "role": "Organization"})
        headers = {"Authorization": f"Bearer {token}"}

        # Create a test customer with potential formula injection payloads
        malicious_cust = Customer(
            name="=SUM(A1:A10)",
            phone="+919999988888",
            email="+malicious@example.com",
            address="@calc.exe",
            organization_id=1,
            notes="-CMD|'/C calc'!A0"
        )
        db.session.add(malicious_cust)
        db.session.commit()

        # 1. Test format=xlsx
        res_xlsx = self.client.get("/api/v1/customers/export?format=xlsx", headers=headers)
        self.assertEqual(res_xlsx.status_code, 200)
        self.assertEqual(res_xlsx.content_type, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
        self.assertIn("customers_export_1.xlsx", res_xlsx.headers.get("Content-Disposition", ""))

        # Verify workbook content and sanitization using openpyxl
        import openpyxl
        wb = openpyxl.load_workbook(io.BytesIO(res_xlsx.get_data()))
        ws = wb.active
        self.assertEqual(ws.title, "Customers")
        
        # Verify headers
        header_vals = [cell.value for cell in ws[1]]
        self.assertIn("Name", header_vals)
        self.assertIn("Phone", header_vals)
        self.assertIn("Email", header_vals)

        # Find our customer row and verify formula injection was prepended with single quote
        found_sanitized = False
        for row in ws.iter_rows(min_row=2, values_only=True):
            if row[1] and row[1].startswith("'=SUM"):
                found_sanitized = True
                self.assertTrue(row[6].startswith("'@calc.exe"))
                self.assertTrue(row[8].startswith("'-CMD"))
                break
        self.assertTrue(found_sanitized, "Malicious cell formula should be sanitized with leading single quote")

        # 2. Test format=csv also sanitizes
        res_csv = self.client.get("/api/v1/customers/export?format=csv", headers=headers)
        self.assertEqual(res_csv.status_code, 200)
        self.assertIn("text/csv", res_csv.content_type)
        csv_data = res_csv.get_data(as_text=True)
        self.assertIn("'=SUM", csv_data)

        # 3. Test unsupported format -> 400
        res_bad = self.client.get("/api/v1/customers/export?format=pdf_not_supported", headers=headers)
        self.assertEqual(res_bad.status_code, 400)

        # Clean up test customer
        db.session.delete(malicious_cust)
        db.session.commit()

    def test_financial_ledger_integrity_and_reconciliation(self):
        """Validates that session credit redemptions produce immutable FinancialLedgerEntry rows with verified reconciliation."""
        token = generate_access_token(2, "Organization", extra_claims={"organization_id": "2", "role": "Organization"})
        headers = {"Authorization": f"Bearer {token}"}

        rec = IndustryRecord.query.filter_by(organization_id=2).first()
        self.assertIsNotNone(rec)

        # Initial state setup
        rec_data = dict(rec.data or {})
        rec_data["total_sessions"] = 12
        rec_data["remaining_sessions"] = 12
        rec.data = rec_data
        db.session.commit()

        # Redeem 3 credits
        payload = {
            "sessions_to_redeem": 3,
            "service_name": "Hydrotherapy Ritual",
            "staff_name": "Chloe Senior Stylist",
            "notes": "Session 1 of package"
        }
        res = self.client.post(f"/api/v1/industry-suite/records/{rec.id}/redeem-session-credit", json=payload, headers=headers)
        self.assertEqual(res.status_code, 200)
        data = res.get_json()
        self.assertEqual(data["remaining_credits"], 9)
        self.assertIn("ledger_entry_id", data)

        # Verify ledger entry directly from database
        entry = FinancialLedgerEntry.query.get(data["ledger_entry_id"])
        self.assertIsNotNone(entry)
        self.assertEqual(entry.entry_type, "CREDIT_REDEMPTION")
        self.assertEqual(entry.units_delta, -3.0)
        self.assertEqual(entry.balance_after, 9.0)
        self.assertEqual(entry.organization_id, 2)

        # Verify ledger-audit endpoint
        res_audit = self.client.get(f"/api/v1/industry-suite/records/{rec.id}/ledger-audit", headers=headers)
        self.assertEqual(res_audit.status_code, 200)
        audit_data = res_audit.get_json()
        self.assertTrue(audit_data["reconciliation_verified"])
        self.assertEqual(audit_data["latest_ledger_balance"], 9.0)
        self.assertGreater(audit_data["entries_count"], 0)

    def test_webhook_cryptographic_signature_verification(self):
        """Validates HMAC-SHA256 signature enforcement on Razorpay webhook endpoint."""
        from app.routes.payments import PAYMENT_CONFIG
        webhook_secret = PAYMENT_CONFIG.get("webhook_secret", "rzp_webhook_secret_secure_2026")
        
        # 1. Reject invalid signature
        payload_dict = {
            "event": "payment.captured",
            "payload": {
                "payment": {
                    "entity": {
                        "id": "pay_test_fixture_123",
                        "amount": 50000,
                        "notes": {"tenant_id": "1"}
                    }
                }
            }
        }
        payload_bytes = json.dumps(payload_dict).encode("utf-8")

        bad_headers = {
            "Content-Type": "application/json",
            "X-Razorpay-Signature": "invalid_tampered_signature_hex"
        }
        res_bad = self.client.post("/api/v1/payments/webhook", data=payload_bytes, headers=bad_headers)
        self.assertEqual(res_bad.status_code, 400)
        self.assertIn("INVALID_SIGNATURE", res_bad.get_data(as_text=True))

        # 2. Accept validly signed signature
        valid_sig = hmac.new(
            webhook_secret.encode("utf-8"),
            payload_bytes,
            hashlib.sha256
        ).hexdigest()

        good_headers = {
            "Content-Type": "application/json",
            "X-Razorpay-Signature": valid_sig
        }
        res_good = self.client.post("/api/v1/payments/webhook", data=payload_bytes, headers=good_headers)
        self.assertEqual(res_good.status_code, 200, res_good.get_data(as_text=True))
        good_json = res_good.get_json()
        self.assertEqual(good_json["status"], "processed")
        self.assertEqual(good_json["payment_id"], "pay_test_fixture_123")

        # 3. Reject webhook when no tenant context can be bound
        payload_orphan = {
            "event": "payment.captured",
            "payload": {
                "payment": {
                    "entity": {
                        "id": "pay_orphan_999",
                        "amount": 25000,
                        "notes": {}
                    }
                }
            }
        }
        orphan_bytes = json.dumps(payload_orphan).encode("utf-8")
        orphan_sig = hmac.new(
            webhook_secret.encode("utf-8"),
            orphan_bytes,
            hashlib.sha256
        ).hexdigest()

        res_orphan = self.client.post("/api/v1/payments/webhook", data=orphan_bytes, headers={"Content-Type": "application/json", "X-Razorpay-Signature": orphan_sig})
        self.assertEqual(res_orphan.status_code, 400)
        self.assertIn("TENANT_REQUIRED", res_orphan.get_data(as_text=True))


if __name__ == "__main__":
    unittest.main()
