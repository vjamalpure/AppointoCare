from datetime import datetime
from flask import Blueprint, request, jsonify
from flask_jwt_extended import jwt_required, get_jwt, get_jwt_identity
from app.models import db, Customer, Organization, Appointment, MessageLog, AppointmentTransaction, AuditLog

customer_bp = Blueprint("customer_bp", __name__)


def _sanitize_cell(val) -> str:
    """Neutralizes formula injection (CWE-1236) in CSV and spreadsheet exports."""
    if val is None:
        return ""
    text = str(val).strip()
    if text and text[0] in ("=", "+", "-", "@", "\t", "\r", "%"):
        return f"'{text}"
    return text


@customer_bp.route("", methods=["GET"])
@customer_bp.route("/all", methods=["GET"])
@jwt_required()
def get_customers():
    claims = get_jwt()
    role = claims.get("role")
    org_id = request.args.get("organization_id")

    query = Customer.query
    if role in ["Admin", "SuperAdmin"]:
        if org_id and org_id != "ALL":
            query = query.filter_by(organization_id=int(org_id))
    else:
        org_id = int(claims.get("organization_id") or 0)
        query = query.filter_by(organization_id=org_id)

    q = request.args.get("q") or request.args.get("search")
    if q:
        query = query.filter(
            db.or_(
                Customer.name.ilike(f"%{q}%"),
                Customer.phone.ilike(f"%{q}%"),
                Customer.email.ilike(f"%{q}%")
            )
        )

    customers = query.order_by(Customer.created_at.desc()).all()

    return jsonify([
        {
            "id": c.id,
            "organization_id": c.organization_id,
            "name": c.name,
            "phone": c.phone,
            "email": c.email,
            "gender": c.gender,
            "dob": c.dob.isoformat() if c.dob else None,
            "address": c.address,
            "tags": c.tags or "General",
            "notes": c.notes,
            "source": c.source or "Direct",
            "created_at": c.created_at.isoformat() if c.created_at else None,
            "updated_at": c.updated_at.isoformat() if c.updated_at else None,
        }
        for c in customers
    ])


@customer_bp.route("", methods=["POST"])
@customer_bp.route("/create", methods=["POST"])
@jwt_required()
def create_customer():
    claims = get_jwt()
    role = claims.get("role")
    data = request.json or {}

    if role in ["Admin", "SuperAdmin"]:
        organization_id = data.get("organization_id") or claims.get("organization_id") or 1
    else:
        organization_id = int(claims.get("organization_id") or 0)

    if not data.get("name") or not data.get("phone"):
        return jsonify({"msg": "name and phone are required"}), 400

    phone = str(data["phone"]).strip()
    check_duplicate = data.get("check_duplicate", False)
    existing = Customer.query.filter_by(organization_id=int(organization_id), phone=phone).first()
    if existing and check_duplicate and not data.get("allow_duplicate", False):
        return jsonify({
            "msg": f"Patient with phone '{phone}' already registered in this organization",
            "existing_customer_id": existing.id,
            "existing_name": existing.name
        }), 409

    customer = Customer(
        organization_id=int(organization_id),
        name=data["name"],
        phone=phone,
        email=data.get("email"),
        gender=data.get("gender", "Female"),
        dob=data.get("dob"),
        address=data.get("address"),
        tags=data.get("tags", "General"),
        notes=data.get("notes"),
        source=data.get("source", "AppointoCare Web Intake")
    )
    db.session.add(customer)

    # Sync with Patient model if not already present
    from app.models import Patient
    patient = Patient.query.filter_by(organization_id=int(organization_id), phone=phone).first()
    if not patient:
        patient = Patient(
            organization_id=int(organization_id),
            name=data["name"],
            phone=phone,
            email=data.get("email"),
            gender=data.get("gender", "Female"),
            address=data.get("address")
        )
        db.session.add(patient)

    db.session.commit()

    return jsonify({
        "msg": "Customer created successfully",
        "customer_id": customer.id,
        "id": customer.id,
        "name": customer.name,
        "phone": customer.phone
    }), 201


@customer_bp.route("/<int:customer_id>", methods=["GET"])
@jwt_required()
def get_customer(customer_id):
    claims = get_jwt()
    role = claims.get("role")
    customer = Customer.query.get_or_404(customer_id)

    if role != "Admin" and customer.organization_id != int(claims.get("organization_id") or 0):
        return jsonify({"msg": "Unauthorized"}), 403

    return jsonify({
        "id": customer.id,
        "organization_id": customer.organization_id,
        "name": customer.name,
        "phone": customer.phone,
        "email": customer.email,
        "gender": customer.gender,
        "dob": customer.dob.isoformat() if customer.dob else None,
        "address": customer.address,
        "tags": customer.tags,
        "notes": customer.notes,
        "source": customer.source,
        "created_at": customer.created_at.isoformat() if customer.created_at else None,
        "updated_at": customer.updated_at.isoformat() if customer.updated_at else None,
    })


@customer_bp.route("/<int:customer_id>", methods=["PUT", "PATCH"])
@jwt_required()
def update_customer(customer_id):
    claims = get_jwt()
    role = claims.get("role")
    data = request.json or {}

    customer = Customer.query.get_or_404(customer_id)
    if role != "Admin" and customer.organization_id != int(claims.get("organization_id") or 0):
        return jsonify({"msg": "Unauthorized"}), 403

    if "name" in data:
        customer.name = data["name"]
    if "phone" in data:
        customer.phone = data["phone"]
    if "email" in data:
        customer.email = data["email"]
    if "gender" in data:
        customer.gender = data["gender"]
    if "dob" in data:
        customer.dob = data["dob"]
    if "address" in data:
        customer.address = data["address"]
    if "tags" in data:
        customer.tags = data["tags"]
    if "notes" in data:
        customer.notes = data["notes"]
    if "source" in data:
        customer.source = data["source"]

    db.session.commit()
    return jsonify({"msg": "Customer updated successfully", "id": customer.id})


@customer_bp.route("/<int:customer_id>", methods=["DELETE"])
@jwt_required()
def delete_customer(customer_id):
    claims = get_jwt()
    role = claims.get("role")
    customer = Customer.query.get_or_404(customer_id)

    if role != "Admin" and customer.organization_id != int(claims.get("organization_id") or 0):
        return jsonify({"msg": "Unauthorized"}), 403

    db.session.delete(customer)
    db.session.commit()
    return jsonify({"msg": "Customer deleted successfully"})


@customer_bp.route("/<int:customer_id>/timeline", methods=["GET"])
@jwt_required()
def get_customer_timeline(customer_id):
    claims = get_jwt()
    role = claims.get("role")
    customer = Customer.query.get_or_404(customer_id)

    if role != "Admin" and customer.organization_id != int(claims.get("organization_id") or 0):
        return jsonify({"msg": "Unauthorized"}), 403

    # Gather appointments
    appts = Appointment.query.filter(
        Appointment.organization_id == customer.organization_id,
        db.or_(Appointment.customer_phone == customer.phone, Appointment.customer_name == customer.name)
    ).order_by(Appointment.appointment_date.desc()).all()

    appt_ids = [a.id for a in appts]

    # Gather messages
    msgs = MessageLog.query.filter(
        MessageLog.organization_id == customer.organization_id,
        MessageLog.recipient_number == customer.phone
    ).order_by(MessageLog.created_at.desc()).all()

    # Gather clinical encounters, prescriptions, and diagnostics from IndustryRecord
    from app.models import IndustryRecord
    records = IndustryRecord.query.filter(
        IndustryRecord.organization_id == customer.organization_id,
        db.or_(
            IndustryRecord.customer_id == customer.id,
            IndustryRecord.appointment_id.in_(appt_ids) if appt_ids else False
        )
    ).order_by(IndustryRecord.created_at.desc()).all()

    # Gather billing transactions
    txns = AppointmentTransaction.query.filter(
        AppointmentTransaction.organization_id == customer.organization_id,
        AppointmentTransaction.appointment_id.in_(appt_ids) if appt_ids else False
    ).order_by(AppointmentTransaction.created_at.desc()).all()

    timeline = []
    for a in appts:
        timeline.append({
            "type": "appointment",
            "title": f"Appointment: {a.status}",
            "description": f"Scheduled on {a.appointment_date.strftime('%Y-%m-%d %H:%M') if a.appointment_date else 'N/A'}. Provider: {a.staff_name or 'General'} | Payment: {a.payment_status}",
            "date": a.appointment_date.isoformat() if a.appointment_date else a.created_at.isoformat(),
            "status": a.status,
            "payment_status": a.payment_status,
            "appointment_id": a.id
        })

    for m in msgs:
        timeline.append({
            "type": "whatsapp",
            "title": f"WhatsApp ({m.status})",
            "description": m.message_content,
            "date": m.created_at.isoformat() if m.created_at else None,
            "status": m.status
        })

    for r in records:
        rec_label = r.record_type.replace('_', ' ').title()
        timeline.append({
            "type": "clinical_record",
            "record_type": r.record_type,
            "title": f"{rec_label}: {r.title}",
            "description": f"Status: {r.status} | Practitioner: {r.created_by_user or 'Specialist'}",
            "date": r.created_at.isoformat() if r.created_at else None,
            "status": r.status,
            "record_id": r.id
        })

    for t in txns:
        timeline.append({
            "type": "billing",
            "title": f"Receipt: ₹{t.amount} ({t.payment_method})",
            "description": f"{t.transaction_type} | Status: {t.status}",
            "date": t.created_at.isoformat() if t.created_at else None,
            "status": t.status,
            "amount": t.amount,
            "transaction_id": t.id
        })

    timeline.sort(key=lambda x: x.get("date") or "", reverse=True)

    return jsonify({
        "customer": {
            "id": customer.id,
            "name": customer.name,
            "phone": customer.phone,
            "email": customer.email,
            "tags": customer.tags,
            "notes": customer.notes,
            "source": customer.source
        },
        "timeline": timeline,
        "total_appointments": len(appts),
        "total_messages": len(msgs),
        "total_clinical_records": len(records),
        "total_transactions": len(txns)
    })


@customer_bp.route("/history", methods=["GET"])
@jwt_required()
def get_customer_history():
    claims = get_jwt()
    role = claims.get("role")
    phone = request.args.get("phone")
    name = request.args.get("name")
    org_id = int(claims.get("organization_id") or 0) if role != "Admin" else request.args.get("organization_id")

    query = Appointment.query
    if org_id:
        query = query.filter(Appointment.organization_id == int(org_id))
    if phone:
        query = query.filter(Appointment.customer_phone.ilike(f"%{phone}%"))
    if name:
        query = query.filter(Appointment.customer_name.ilike(f"%{name}%"))

    appts = query.order_by(Appointment.appointment_date.desc()).limit(20).all()

    return jsonify([
        {
            "id": a.id,
            "customer_name": a.customer_name,
            "customer_phone": a.customer_phone,
            "appointment_date": a.appointment_date.isoformat() if a.appointment_date else None,
            "status": a.status,
            "payment_status": a.payment_status,
            "organization_id": a.organization_id
        }
        for a in appts
    ])


@customer_bp.route("/export", methods=["GET"])
@jwt_required()
def export_customers():
    claims = get_jwt()
    role = claims.get("role")
    
    query = Customer.query
    if role in ["Admin", "SuperAdmin"]:
        org_id = request.args.get("organization_id")
        if org_id and org_id != "ALL":
            query = query.filter_by(organization_id=int(org_id))
    else:
        scoped_org_id = int(claims.get("organization_id") or 0)
        query = query.filter_by(organization_id=scoped_org_id)

    tag = request.args.get("tag")
    if tag:
        query = query.filter(Customer.tags.ilike(f"%{tag}%"))

    q = request.args.get("q") or request.args.get("search")
    if q:
        query = query.filter(
            db.or_(
                Customer.name.ilike(f"%{q}%"),
                Customer.phone.ilike(f"%{q}%"),
                Customer.email.ilike(f"%{q}%")
            )
        )

    customers = query.order_by(Customer.created_at.desc()).all()

    export_format = (request.args.get("format") or request.args.get("export_format") or "json").lower()
    if export_format not in ["json", "csv", "xlsx"]:
        return jsonify({"msg": f"Unsupported export format '{export_format}'. Allowed: json, csv, xlsx"}), 400

    def _sanitize_cell(val):
        if val is None:
            return ""
        s = str(val)
        if s and s[0] in ("=", "+", "-", "@", "\t", "\r"):
            return f"'{s}"
        return s

    # Record audit log
    try:
        user_id_val = claims.get("sub")
        org_id_for_log = int(claims.get("organization_id") or 0)
        audit = AuditLog(
            organization_id=org_id_for_log,
            user_id=int(user_id_val) if user_id_val and str(user_id_val).isdigit() else None,
            user_role=role,
            action="EXPORT_CUSTOMERS",
            entity="Customer",
            details=f"Exported {len(customers)} customer records (format: {export_format}, tag filter: {tag or 'None'}, search: {q or 'None'})",
            ip_address=request.remote_addr
        )
        db.session.add(audit)
        db.session.commit()
    except Exception:
        db.session.rollback()

    org_tag = scoped_org_id if 'scoped_org_id' in locals() else (org_id if 'org_id' in locals() else "all")

    if export_format == "csv":
        import io
        import csv
        from flask import Response
        output = io.StringIO()
        writer = csv.writer(output)
        writer.writerow(["ID", "Name", "Phone", "Email", "Gender", "DOB", "Address", "Tags", "Notes", "Source", "Created At"])
        for c in customers:
            writer.writerow([
                c.id,
                _sanitize_cell(c.name),
                _sanitize_cell(c.phone),
                _sanitize_cell(c.email),
                _sanitize_cell(c.gender),
                c.dob.isoformat() if c.dob else "",
                _sanitize_cell(c.address),
                _sanitize_cell(c.tags),
                _sanitize_cell(c.notes),
                _sanitize_cell(c.source),
                c.created_at.isoformat() if c.created_at else ""
            ])
        output.seek(0)
        return Response(
            output.getvalue(),
            mimetype="text/csv",
            headers={"Content-Disposition": f"attachment; filename=customers_export_{org_tag}.csv"}
        )

    elif export_format == "xlsx":
        import io
        from flask import Response
        import openpyxl
        from openpyxl.styles import Font, PatternFill, Alignment

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Customers"

        headers = ["ID", "Name", "Phone", "Email", "Gender", "DOB", "Address", "Tags", "Notes", "Source", "Created At"]
        ws.append(headers)

        header_fill = PatternFill(start_color="1E3A8A", end_color="1E3A8A", fill_type="solid")
        header_font = Font(name="Arial", size=11, bold=True, color="FFFFFF")
        for col_num in range(1, len(headers) + 1):
            cell = ws.cell(row=1, column=col_num)
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center", vertical="center")

        for c in customers:
            ws.append([
                c.id,
                _sanitize_cell(c.name),
                _sanitize_cell(c.phone),
                _sanitize_cell(c.email),
                _sanitize_cell(c.gender),
                c.dob.isoformat() if c.dob else "",
                _sanitize_cell(c.address),
                _sanitize_cell(c.tags),
                _sanitize_cell(c.notes),
                _sanitize_cell(c.source),
                c.created_at.isoformat() if c.created_at else ""
            ])

        for col in ws.columns:
            max_len = max(len(str(cell.value or "")) for cell in col)
            col_letter = openpyxl.utils.get_column_letter(col[0].column)
            ws.column_dimensions[col_letter].width = min(max(max_len + 3, 12), 40)

        output = io.BytesIO()
        wb.save(output)
        output.seek(0)
        return Response(
            output.getvalue(),
            mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            headers={"Content-Disposition": f"attachment; filename=customers_export_{org_tag}.xlsx"}
        )

    return jsonify({
        "count": len(customers),
        "exported_at": datetime.utcnow().isoformat(),
        "format": "json",
        "customers": [
            {
                "id": c.id,
                "organization_id": c.organization_id,
                "name": c.name,
                "phone": c.phone,
                "email": c.email,
                "gender": c.gender,
                "dob": c.dob.isoformat() if c.dob else None,
                "address": c.address,
                "tags": c.tags,
                "notes": c.notes,
                "source": c.source,
                "created_at": c.created_at.isoformat() if c.created_at else None
            }
            for c in customers
        ]
    })


# -----------------------------------------------------------------------------
# Customer 360 Unified Communication Feed
# -----------------------------------------------------------------------------
@customer_bp.route("/<int:customer_id>/unified-timeline", methods=["GET"])
@customer_bp.route("/<int:customer_id>/360", methods=["GET"])
@jwt_required()
def get_customer_unified_timeline(customer_id):
    claims = get_jwt()
    role = claims.get("role")
    customer = Customer.query.get_or_404(customer_id)

    if role not in ["Admin", "SuperAdmin"]:
        org_id = int(claims.get("organization_id") or 0)
        if customer.organization_id != org_id:
            return jsonify({"msg": "Unauthorized"}), 403

    from app.models import (
        ClinicalEncounter, CustomerDocument, TaskItem, FinancialLedgerEntry,
        BespokeOrder, InsurancePolicyRecord, LegalMatter, StudentReadinessScore
    )

    events = []

    # 1. Appointments
    appts = Appointment.query.filter_by(customer_id=customer.id, organization_id=customer.organization_id).all()
    for a in appts:
        events.append({
            "type": "APPOINTMENT",
            "title": f"Appointment: {a.service_name or 'General Visit'}",
            "description": f"Status: {a.status} with {a.staff_name or 'Staff'}. Notes: {a.notes or 'None'}",
            "timestamp": a.appointment_date.isoformat(),
            "badge_color": "blue" if a.status == "Booked" else ("green" if a.status == "Completed" else "red")
        })

    # 2. Messages (WhatsApp, SMS, Email)
    msgs = MessageLog.query.filter(
        MessageLog.organization_id == customer.organization_id,
        MessageLog.recipient_number.like(f"%{customer.phone[-10:]}%")
    ).all()
    for m in msgs:
        events.append({
            "type": "COMMUNICATION",
            "title": f"{m.message_type} ({m.direction})",
            "description": m.message_content[:150] + "..." if len(m.message_content) > 150 else m.message_content,
            "timestamp": m.sent_at.isoformat() if m.sent_at else (m.created_at.isoformat() if m.created_at else datetime.utcnow().isoformat()),
            "badge_color": "green" if m.status == "Sent" else "amber"
        })

    # 3. Clinical Encounters (Role-guarded: Clinical credentials or Admin)
    if role in ("Doctor", "Nurse", "Specialist", "Admin", "SuperAdmin", "Organization", "Manager"):
        encounters = ClinicalEncounter.query.filter_by(patient_id=customer.id, organization_id=customer.organization_id).all()
        for e in encounters:
            events.append({
                "type": "CLINICAL_ENCOUNTER",
                "title": f"Encounter: {e.encounter_type} (Dr. {e.practitioner_name})",
                "description": f"Chief Complaints: {e.chief_complaints}. Diagnosis: {e.diagnosis}",
                "timestamp": e.encounter_date.isoformat(),
                "badge_color": "purple"
            })

    # 4. Uploaded Documents
    docs = CustomerDocument.query.filter_by(customer_id=customer.id, organization_id=customer.organization_id).all()
    for d in docs:
        events.append({
            "type": "DOCUMENT",
            "title": f"Document: {d.title} ({d.file_name})",
            "description": f"Uploaded by {d.uploaded_by}. Status: {d.status}",
            "timestamp": d.created_at.isoformat() if d.created_at else datetime.utcnow().isoformat(),
            "badge_color": "teal"
        })

    # 5. Financial Ledger & Transactions
    ledgers = FinancialLedgerEntry.query.filter_by(customer_id=customer.id, organization_id=customer.organization_id).all()
    for l in ledgers:
        events.append({
            "type": "FINANCIAL",
            "title": f"Ledger: {l.entry_type}",
            "description": f"Delta: {l.units_delta} {l.currency_or_unit}. Balance: {l.balance_after}. {l.description or ''}",
            "timestamp": l.created_at.isoformat() if l.created_at else datetime.utcnow().isoformat(),
            "badge_color": "indigo"
        })

    # 6. Assigned Tasks
    tasks = TaskItem.query.filter_by(customer_id=customer.id, organization_id=customer.organization_id).all()
    for t in tasks:
        events.append({
            "type": "TASK",
            "title": f"Task: {t.title} ({t.priority})",
            "description": f"Assigned to: {t.assigned_to_name or 'Staff'}. Status: {t.status}",
            "timestamp": t.created_at.isoformat() if t.created_at else datetime.utcnow().isoformat(),
            "badge_color": "amber" if t.status != "Completed" else "green"
        })

    # 7. Vertical Extensions (Bespoke Orders, Insurance Policies, Legal Matters)
    bespoke_orders = BespokeOrder.query.filter_by(customer_id=customer.id, organization_id=customer.organization_id).all()
    for bo in bespoke_orders:
        events.append({
            "type": "BESPOKE_ORDER",
            "title": f"Bespoke: {bo.garment_type} ({bo.order_number})",
            "description": f"Stage: {bo.production_stage}. Balance Due: ₹{bo.balance_due}. Fabric: {bo.fabric_choice}",
            "timestamp": bo.created_at.isoformat() if bo.created_at else datetime.utcnow().isoformat(),
            "badge_color": "cyan"
        })

    policies = InsurancePolicyRecord.query.filter_by(customer_id=customer.id, organization_id=customer.organization_id).all()
    for pol in policies:
        events.append({
            "type": "INSURANCE_POLICY",
            "title": f"Policy: {pol.policy_number} ({pol.carrier_name})",
            "description": f"Plan: {pol.plan_type}. Annual Premium: ₹{pol.annual_premium}. Expiry: {pol.expiry_date.strftime('%d %b %Y')}",
            "timestamp": pol.created_at.isoformat() if pol.created_at else datetime.utcnow().isoformat(),
            "badge_color": "emerald"
        })

    matters = LegalMatter.query.filter_by(client_id=customer.id, organization_id=customer.organization_id).all()
    for mat in matters:
        events.append({
            "type": "LEGAL_MATTER",
            "title": f"Legal Matter: {mat.matter_number} - {mat.title}",
            "description": f"Practice: {mat.practice_area}. Retainer Balance: ₹{mat.retainer_balance}. Billable Hours: {mat.billable_hours_recorded}",
            "timestamp": mat.created_at.isoformat() if mat.created_at else datetime.utcnow().isoformat(),
            "badge_color": "amber"
        })

    # Sort descending by timestamp
    events.sort(key=lambda x: x["timestamp"], reverse=True)

    return jsonify({
        "customer": {
            "id": customer.id,
            "name": customer.name,
            "phone": customer.phone,
            "email": customer.email,
            "tags": customer.tags
        },
        "total_events": len(events),
        "events": events,
        "timeline_events": events
    }), 200


# -----------------------------------------------------------------------------
# Batch CSV/XLSX Import Validation & Duplicate Detection
# -----------------------------------------------------------------------------
@customer_bp.route("/import-validate", methods=["POST"])
@jwt_required()
def validate_customer_import():
    claims = get_jwt()
    role = claims.get("role")
    data = request.json or {}

    if role in ["Admin", "SuperAdmin"]:
        org_id = int(data.get("organization_id") or claims.get("organization_id") or 1)
    else:
        org_id = int(claims.get("organization_id") or 0)

    rows = data.get("rows", [])
    if not rows or not isinstance(rows, list):
        return jsonify({"msg": "rows must be a non-empty list of client records"}), 400

    existing_customers = Customer.query.filter_by(organization_id=org_id).all()
    existing_phones = {c.phone.strip() for c in existing_customers if c.phone}
    existing_emails = {c.email.strip().lower() for c in existing_customers if c.email}

    valid_records = []
    duplicate_records = []
    error_records = []

    seen_batch_phones = set()

    for idx, row in enumerate(rows, start=1):
        name = str(row.get("name", "")).strip()
        phone = str(row.get("phone", "")).strip()
        email = str(row.get("email", "")).strip().lower()

        if not name:
            error_records.append({"row": idx, "reason": "Missing required field 'name'", "data": row})
            continue

        if not phone or len(phone) < 7:
            error_records.append({"row": idx, "reason": "Invalid or missing 'phone'", "data": row})
            continue

        # Check existing in org DB or duplicate in batch
        if phone in existing_phones or (email and email in existing_emails):
            duplicate_records.append({"row": idx, "reason": "Already exists in database", "data": row})
            continue

        if phone in seen_batch_phones:
            duplicate_records.append({"row": idx, "reason": "Duplicate phone within uploaded file", "data": row})
            continue

        seen_batch_phones.add(phone)
        valid_records.append({
            "name": name,
            "phone": phone,
            "email": email or None,
            "gender": row.get("gender", "Female"),
            "tags": row.get("tags", "Imported Client"),
            "notes": row.get("notes", "Batch CSV/XLSX intake"),
            "source": row.get("source", "Spreadsheet Import")
        })

    return jsonify({
        "summary": {
            "total_rows_scanned": len(rows),
            "valid_count": len(valid_records),
            "duplicate_count": len(duplicate_records),
            "error_count": len(error_records),
            "ready_for_import": len(valid_records) > 0
        },
        "valid_preview": valid_records[:10],
        "duplicates": duplicate_records,
        "errors": error_records
    }), 200


@customer_bp.route("/bulk-import", methods=["POST"])
@jwt_required()
def bulk_import_customers():
    claims = get_jwt()
    role = claims.get("role")
    data = request.json or {}

    if role in ["Admin", "SuperAdmin"]:
        org_id = int(data.get("organization_id") or claims.get("organization_id") or 1)
    else:
        org_id = int(claims.get("organization_id") or 0)

    records = data.get("records", [])
    if not records:
        return jsonify({"msg": "No records supplied for import"}), 400

    imported_count = 0
    for r in records:
        cust = Customer(
            organization_id=org_id,
            name=r["name"],
            phone=r["phone"],
            email=r.get("email"),
            gender=r.get("gender", "Female"),
            tags=r.get("tags", "Imported"),
            notes=r.get("notes"),
            source=r.get("source", "Batch Import")
        )
        db.session.add(cust)
        imported_count += 1

    # Log audit
    audit = AuditLog(
        organization_id=org_id,
        action=f"BULK_CUSTOMER_IMPORT: {imported_count} records",
        details=f"Bulk imported by {claims.get('sub') or 'Staff'}",
        user_role=claims.get("role") or "Staff",
        ip_address=request.remote_addr or "127.0.0.1"
    )
    db.session.add(audit)
    db.session.commit()

    return jsonify({
        "msg": f"Successfully imported {imported_count} customer records",
        "imported_count": imported_count
    }), 201


@customer_bp.route("/<int:customer_id>/portal-link", methods=["POST"])
@jwt_required()
def generate_customer_portal_link(customer_id):
    """Generates a secure, 72-hour customer self-service portal magic link."""
    claims = get_jwt()
    role = claims.get("role")
    org_id = int(claims.get("organization_id") or 0) if role not in ["Admin", "SuperAdmin"] else None

    query = Customer.query.filter_by(id=customer_id)
    if org_id:
        query = query.filter_by(organization_id=org_id)
    customer = query.first_or_404()

    import secrets
    from datetime import timedelta
    token = secrets.token_hex(32)
    auth_code = f"{secrets.randbelow(900000) + 100000}"
    expires = datetime.utcnow() + timedelta(hours=72)

    from app.models import CustomerPortalSession
    # Invalidate previous sessions
    CustomerPortalSession.query.filter_by(customer_id=customer.id, is_active=True).update({"is_active": False})

    session = CustomerPortalSession(
        customer_id=customer.id,
        organization_id=customer.organization_id,
        session_token=token,
        auth_code=auth_code,
        expires_at=expires,
        is_active=True,
        communication_preferences={"whatsapp": True, "email": True}
    )
    db.session.add(session)
    db.session.commit()

    portal_url = f"/portal/login?token={token}&org={customer.organization_id}"
    return jsonify({
        "msg": "Customer portal magic link generated",
        "customer_id": customer.id,
        "customer_name": customer.name,
        "session_token": token,
        "auth_code": auth_code,
        "portal_url": portal_url,
        "expires_at": expires.isoformat()
    }), 200



